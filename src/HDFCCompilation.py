import openpyxl
from openpyxl import Workbook
import os
import re
import datetime


LOG_FILE = "bank4_log.txt"

TARGETS = {
    "grand_total": "grand total",
    "net_assets": "net current assets",
    "isin": "isin"
}

# =========================
# LOGGING
# =========================
def log(msg):
    t = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = msg.encode("ascii", "replace").decode("ascii")
    print(f"{t} {safe}")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{t} {safe}\n")

# =========================
# CLEAN
# =========================
def clean(v):
    if v is None:
        return ""
    v = str(v).lower()
    v = v.replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", v).strip()

# =========================
# TREPS / REVERSE REPO CHECK
# =========================
def is_treps_or_reverse_repo(text: str) -> bool:
    t = clean(text)
    return ("treps" in t) or ("reverse repo" in t)

# =========================
# PROCESS ONE FILE
# =========================
def process_single_file(filepath, out_ws):
    log(f"\n=== Processing file: {filepath} ===")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # Header injection from A1
    header = ws["A1"].value
    log(f"Header extracted: {header}")

    max_row = ws.max_row
    max_col = ws.max_column

    # Find GRAND TOTAL in column B
    grand_total_row = None
    for r in range(2, max_row + 1):
        if clean(ws[f"B{r}"].value) == TARGETS["grand_total"]:
            grand_total_row = r
            log(f"Found GRAND TOTAL at row {r}")
            break

    if not grand_total_row:
        log("No GRAND TOTAL found — skipping file")
        return

    r = 2
    while r <= grand_total_row:

        colB = clean(ws[f"B{r}"].value)
        colD = clean(ws[f"D{r}"].value)

        # Inject header into column A
        ws[f"A{r}"].value = header

        # =========================
        # COLUMN B HAS DATA (ISIN)
        # =========================
        if colB != "":
            if colB == TARGETS["isin"]:
                r += 1
                continue

            # ISIN must start with IN
            if not colB.startswith("in"):
                r += 1
                continue

            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, max_col + 1)
            ])
            log(f"Row {r}: ISIN row copied")
            r += 1
            continue

        # =========================
        # COLUMN B BLANK → CHECK D
        # =========================
        if colD == "":
            r += 1
            continue

        if colD == TARGETS["isin"]:
            r += 1
            continue

        # Net Current Assets
        if colD == TARGETS["net_assets"]:
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, max_col + 1)
            ])
            log(f"Row {r}: Net Current Assets copied")
            r += 1
            continue

        # 🔥 TREPS / REVERSE REPO (FIXED LOGIC)
        if is_treps_or_reverse_repo(colD):
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, max_col + 1)
            ])
            log(f"Row {r}: TREPS / Reverse Repo copied")
            r += 1
            continue

        r += 1

# =========================
# PROCESS FOLDER
# =========================
def process_folder(folder_path, output_file):
    open(LOG_FILE, "w").close()
    log("=== BANK #4 SCRIPT STARTED ===")

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    files_processed = 0

    for filename in os.listdir(folder_path):
        if filename.startswith("~$"):
            continue
        if filename.endswith(".xlsx"):
            process_single_file(os.path.join(folder_path, filename), out_ws)
            files_processed += 1

    out_wb.save(output_file)

    log("===== SCRIPT COMPLETED =====")
    log(f"TOTAL FILES PROCESSED: {files_processed}")
    log(f"Output saved: {output_file}")

# =========================
# RUN
# =========================
if __name__ == "__main__":
    process_folder(
        folder_path="HDFC",
        output_file="./CompiledFiles/Compiled_HDFC.xlsx"
    )
