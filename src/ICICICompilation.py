import openpyxl
from openpyxl import Workbook
import os
import re
import datetime


# =========================
# BANK #5 TARGET STRINGS
# =========================

TARGETS = {
    "net_assets": "net current assets",
    "stop_key": "total net assets",
    "isin": "isin"
}

LOG_FILE = "bank5_log.txt"


# =========================
# LOGGING
# =========================

def log(msg):
    t = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = msg.encode("ascii", "replace").decode("ascii")
    print(f"{t} {safe}")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{t} {safe}\n")


# =========================
# CLEANING FUNCTION
# =========================

def clean(v):
    if v is None:
        return ""
    v = str(v).lower()
    v = v.replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", v).strip()


# =========================
# TREPS / REVERSE REPO CHECK
# =========================

def is_treps_or_reverse_repo(text: str) -> bool:
    t = clean(text)
    return (
        "treps" in t
        or "reverse repo" in t
    )


# =========================
# PROCESS A SINGLE FILE
# =========================

def process_single_file(filepath, out_ws):
    log(f"\n=== Processing file: {filepath} ===")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active  # Single-sheet input

    # ------------------------------------------
    # HEADER injection: B2 → Column A
    # ------------------------------------------
    header = ws["B2"].value
    log(f"Header extracted from B2: {header}")

    r = 2
    while r <= ws.max_row:

        colB = clean(ws[f"B{r}"].value)
        colC = clean(ws[f"C{r}"].value)

        # Inject header into Column A
        ws[f"A{r}"].value = header

        # ================================
        # RULE 1: COLUMN C HAS DATA (ISIN)
        # ================================
        if colC != "":
            if colC == TARGETS["isin"]:
                log(f"Row {r}: C = ISIN → SKIPPED")
                r += 1
                continue

            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)
            ])
            log(f"Row {r}: ISIN in C → COPIED")
            r += 1
            continue

        # ================================
        # RULE 2: COLUMN C BLANK → CHECK B
        # ================================

        # STOP condition
        if colB == TARGETS["stop_key"]:
            log(f"Row {r}: Found Total Net Assets → STOP FILE")
            break

        # TREPS / REVERSE REPO → COPY
        if is_treps_or_reverse_repo(colB):
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)
            ])
            log(f"Row {r}: TREPS / Reverse Repo → COPIED")
            r += 1
            continue

        # Net Current Assets → COPY
        if colB == TARGETS["net_assets"]:
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)
            ])
            log(f"Row {r}: Net Current Assets → COPIED")
            r += 1
            continue

        # Skip ISIN text in B
        if colB == TARGETS["isin"]:
            log(f"Row {r}: B = ISIN → SKIPPED")
            r += 1
            continue

        log(f"Row {r}: Random B data → SKIPPED")
        r += 1


# =========================
# PROCESS FOLDER
# =========================

def process_folder(folder_path, output_file):
    open(LOG_FILE, "w").close()
    log("=== BANK #5 SCRIPT STARTED ===")

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    files_processed = 0

    for filename in os.listdir(folder_path):
        if filename.startswith("~$"):
            continue
        if filename.endswith(".xlsx"):
            process_single_file(os.path.join(folder_path, filename), out_ws)
            files_processed += 1

    out_wb.save(output_file)

    log("===== SCRIPT COMPLETED =====")
    log(f"TOTAL FILES PROCESSED: {files_processed}")
    log(f"Output saved: {output_file}")


# =========================
# RUN SCRIPT
# =========================

if __name__ == "__main__":
    process_folder(
        folder_path="ICICI",
        output_file="./CompiledFiles/Compiled_ICICI.xlsx"
    )
