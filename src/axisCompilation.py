import openpyxl
from openpyxl import Workbook
from rapidfuzz import fuzz
import traceback
import datetime
import re

# =========================
# TARGET PHRASES (Bank #1)
# =========================
TARGETS = {
    "grand_total": "grand total",
    "subtotal": "sub total",
    "isin": "isin",
    "net_receivables": "net receivables / (payables)",
    "treps": "reverse repo / treps"     # <<-- CHANGED FOR BANK #1
}

STRICT_FUZZ = 85   # For SUB TOTAL & ISIN
SOFT_FUZZ = 70     # For TREPS & Net Receivables
LOG_FILE = "script_log.txt"


# =========================
# LOGGING
# =========================
def log(message: str) -> None:
    timestamp = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = message.encode("ascii", "replace").decode("ascii")
    line = f"{timestamp} {safe}"
    print(line)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")


# =========================
# CLEAN TEXT
# =========================
def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).lower()
    text = text.replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


# =========================
# FUZZY MATCHER
# =========================
def fuzzy_match(value, key: str) -> bool:
    cleaned = clean(value)
    target = TARGETS[key]

    threshold = SOFT_FUZZ if key in ("treps", "net_receivables") else STRICT_FUZZ
    score = fuzz.partial_ratio(cleaned, target)

    log(f"      fuzzy-check: '{cleaned}' vs '{target}' | score={score} threshold={threshold}")
    return score >= threshold


# =========================
# EXACT MATCH FOR GRAND TOTAL ONLY
# =========================
def exact_grand_total(value) -> bool:
    cleaned = clean(value)
    match = cleaned == TARGETS["grand_total"]
    log(f"      exact-check: '{cleaned}' == 'grand total'? -> {match}")
    return match


# =========================
# MAIN ENGINE
# =========================
def process_file(input_path: str, output_path: str) -> None:
    open(LOG_FILE, "w").close()
    log("========== SCRIPT STARTED ==========")

    try:
        wb = openpyxl.load_workbook(input_path, data_only=True)
        log(f"Loaded workbook: {input_path}")
    except Exception as e:
        log(f"ERROR loading workbook: {e}")
        log(traceback.format_exc())
        return

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    for sheet_name in wb.sheetnames:
        log("")
        log(f"--- Processing Sheet: {sheet_name} ---")
        ws = wb[sheet_name]

        # ==========================================
        # FIND GRAND TOTAL EXACTLY
        # ==========================================
        grand_total_row = None

        for r in range(1, ws.max_row + 1):
            if exact_grand_total(ws[f"B{r}"].value):
                grand_total_row = r
                log(f"Found EXACT GRAND TOTAL at row {r}")
                break

        if grand_total_row is None:
            log("No EXACT GRAND TOTAL found. Skipping sheet.")
            continue

        # ==========================================
        # FILL B1 into COLUMN A (Bank #1 change)
        # ==========================================
        header = ws["B1"].value
        log(f"Inserting B1 into column A: {header}")

        for r in range(2, grand_total_row + 1):
            ws[f"A{r}"].value = header      # <<-- CHANGED TO COLUMN A

        # ==========================================
        # PROCESS ROWS UNTIL GRAND TOTAL (inclusive)
        # ==========================================
        r = 2
        while r <= grand_total_row:

            colB_raw = ws[f"B{r}"].value
            colC_raw = ws[f"C{r}"].value
            colB = clean(colB_raw)
            colC = clean(colC_raw)

            log(f"Row {r}: B='{colB}' C='{colC}'")

            # ---------------------------------------------
            # CASE A: Column C has data
            # ---------------------------------------------
            if colC != "":
                if fuzzy_match(colC, "isin"):
                    log(f"Row {r}: ISIN detected -> skipping")
                    r += 1
                    continue

                out_ws.append([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])
                log(f"Row {r}: Copied (C has data)")
                r += 1
                continue

            # ---------------------------------------------
            # CASE B: Column C blank
            # ---------------------------------------------
            if fuzzy_match(colB, "net_receivables"):
                out_ws.append([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])
                log(f"Row {r}: Copied (Net Receivables)")
                r += 1
                continue

            if fuzzy_match(colB, "treps"):    # <<-- CHANGED TARGET
                log(f"Row {r}: TREPS detected -> scanning block")
                r += 1  # skip the TREPS row

                while r <= grand_total_row:
                    b_clean = clean(ws[f"B{r}"].value)

                    # STOP before SUB TOTAL
                    if fuzzy_match(b_clean, "subtotal"):
                        log(f"Row {r}: SUB TOTAL found -> ending TREPS block")
                        break

                    out_ws.append([
                        ws.cell(row=r, column=c).value
                        for c in range(1, ws.max_column + 1)
                    ])
                    log(f"Row {r}: Copied (TREPS block)")
                    r += 1

                continue

            r += 1  # default step

        log(f"Finished sheet: {sheet_name}")

    # Save output
    try:
        out_wb.save(output_path)
        log(f"Saved output -> {output_path}")
    except Exception as e:
        log(f"ERROR saving output: {e}")
        log(traceback.format_exc())

    log("========== SCRIPT COMPLETE ==========")


# =========================
# RUN SCRIPT
# =========================
if __name__ == "__main__":
    process_file(
        input_path="Axis.xlsx",
        output_path="./CompiledFiles/Compiled_AxisBank.xlsx"
    )
