import openpyxl
from openpyxl import Workbook
import os
import re
import datetime


# =========================
# BANK #3 TARGET KEYWORDS
# =========================

TARGETS = {
    "grand_total": "grand total",
    "treps": "treps / reverse repo instrument",
    "treps_end": "total",                # TREPS block ends at “Total”
    "net_assets": "net current assets",
    "cash_margin": "cash margin - ccil",
    "isin": "isin"
}

LOG_FILE = "bank3_log.txt"


# =========================
# LOGGING
# =========================

def log(msg):
    t = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = msg.encode("ascii", "replace").decode("ascii")
    print(f"{t} {safe}")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{t} {safe}\n")


# =========================
# CLEANING FUNCTION
# =========================

def clean(v):
    if v is None:
        return ""
    v = str(v).lower()
    v = v.replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", v).strip()


# =========================
# PROCESS A SINGLE FILE
# =========================

def process_single_file(path, out_ws):
    log(f"\n=== Processing file: {path} ===")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active  # Single-sheet files

    # ------------------------------------------
    # HEADER: Row 2, merged across B2:G2 → take B2
    # ------------------------------------------
    header_value = ws["B2"].value
    log(f"Header extracted: {header_value}")

    # ------------------------------------------
    # FIND GRAND TOTAL IN COLUMN B
    # ------------------------------------------
    grand_total_row = None

    for r in range(2, ws.max_row + 1):
        if clean(ws[f"B{r}"].value) == TARGETS["grand_total"]:
            grand_total_row = r
            log(f"Found GRAND TOTAL at row {r}")
            break

    if not grand_total_row:
        log("NO GRAND TOTAL FOUND – SKIPPING FILE")
        return

    # ------------------------------------------
    # PROCESS ROWS
    # ------------------------------------------
    r = 2
    while r <= grand_total_row:

        colB_raw = ws[f"B{r}"].value
        colC_raw = ws[f"C{r}"].value

        colB = clean(colB_raw)
        colC = clean(colC_raw)

        # Inject header into Column A
        ws[f"A{r}"].value = header_value

        # =============================
        # RULE 1: Column C HAS DATA
        # =============================
        if colC != "":

            # Skip ISIN header rows
            if colC == TARGETS["isin"]:
                log(f"Row {r}: Column C = ISIN → SKIPPED")
                r += 1
                continue

            # Otherwise copy entire row
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)
            ])
            log(f"Row {r}: C has data → COPIED")
            r += 1
            continue

        # =============================
        # RULE 2: Column C BLANK → CHECK B
        # =============================

        # If B also blank → skip
        if colB == "":
            log(f"Row {r}: B and C blank → SKIPPED")
            r += 1
            continue

        # Net Current Assets
        if colB == TARGETS["net_assets"]:
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)
            ])
            log(f"Row {r}: Net Current Assets → COPIED")
            r += 1
            continue

        # Cash Margin
        if colB == TARGETS["cash_margin"]:
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)
            ])
            log(f"Row {r}: Cash Margin - CCIL → COPIED")
            r += 1
            continue

        # TREPS trigger
        if colB == TARGETS["treps"]:
            log(f"Row {r}: TREPS trigger → BEGIN BLOCK")
            r += 1

            while r <= grand_total_row:
                nextB = clean(ws[f"B{r}"].value)

                # TREPS ends at "Total"
                if nextB == TARGETS["treps_end"]:
                    log(f"Row {r}: TOTAL → END TREPS BLOCK")
                    break

                out_ws.append([
                    ws.cell(row=r, column=c).value
                    for c in range(1, ws.max_column + 1)
                ])
                log(f"Row {r}: TREPS BLOCK → COPIED")
                r += 1

            continue

        # Anything else in B → skip
        log(f"Row {r}: Random Column B data → SKIPPED")
        r += 1


# =========================
# FOLDER-BASED PROCESSOR
# =========================

def process_folder(folder_path, output_file):
    open(LOG_FILE, "w").close()
    log("=== BANK #3 SCRIPT STARTED ===")

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    for filename in os.listdir(folder_path):
        if filename.startswith("~$"):
            log(f"Skipping temp file: {filename}")
            continue
        if filename.endswith(".xlsx"):
            process_single_file(os.path.join(folder_path, filename), out_ws)


    out_wb.save(output_file)
    log(f"\nSaved output to {output_file}")
    log("=== BANK #3 SCRIPT DONE ===")


# =========================
# RUN SCRIPT
# =========================

if __name__ == "__main__":
    process_folder(
        folder_path="Bandhan",
        output_file="./CompiledFiles/Compiled_BandhanBank.xlsx"
    )
