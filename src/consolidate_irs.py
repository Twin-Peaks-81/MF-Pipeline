import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings
warnings.filterwarnings('ignore')

INPUT_DIR = r'C:\Users\Vedportpc\Documents\Innover\15th May\IRS\final_IRS' + '\\'
OUTPUT_PATH = r'C:\Users\Vedportpc\Documents\Innover\15th May\IRS\final_IRS\IRS_Compilation.xlsx'

DATE_VAL = pd.Timestamp('2026-05-15')

OUT_COLS = [
    'Date', 'MF', 'Fund Name', 'Security Type', 'Name of the Instrument',
    'ISIN', 'Rating Full', 'Rating', 'Quantity', 'Market (Rs. cr)',
    '% to Net\n Assets', 'YTM', 'Maturity Date (Rating)', 'Maturity Date',
    'Maturity(days)', 'Duration', 'Issuer', 'Industry', 'Sector', 'Rating 1',
    'sumproduct yield', 'sumproducct Maturity', 'sumproduct Duration',
    'Month', 'Qtr', 'Year', 'Fortnight', 'Yield Bucket (Liquid)',
    'Yield Bucket', 'Long term Rating', 'Scheme Type', 'Coupon',
    'Perp/Non-Perp', 'Floater details', 'Yr-Month',
    'Swap Type', 'Long Position', 'Short Positon'
]

FIXED = {
    'Date': DATE_VAL,
    'MF': None,
    'Security Type': 'IRS',
    'Rating Full': 'IRS',
    'Rating': 'IRS',
    'Issuer': 'IRS',
    'Industry': 'IRS',
    'Sector': 'IRS',
}

def swap_fields(market_val):
    try:
        v = float(market_val)
    except (TypeError, ValueError):
        return None, None, None
    if v > 0:
        return 'Float to Fixed', 'Receiving Fixed', 'Pay Floating'
    else:
        return 'Fixed to Float', 'Receiving Floating', 'Pay Fixed'

def make_rows(records):
    rows = []
    for rec in records:
        row = {c: None for c in OUT_COLS}
        row.update(FIXED)
        row.update(rec)
        st, lp, sp = swap_fields(row.get('Market (Rs. cr)'))
        row['Swap Type'] = st
        row['Long Position'] = lp
        row['Short Positon'] = sp
        rows.append(row)
    return rows

all_rows = []

# ── AXIS ──────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'Axis_IRS_Output.xlsx', header=0)
for _, r in df.iterrows():
    if pd.isna(r['Instrument Name']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Instrument Name'],
        'Quantity': r['Quantity'],
        'Market (Rs. cr)': r['Market Value'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── BANDHAN ───────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'Bandhan_IRS_Output.xlsx', header=0)
for _, r in df.iterrows():
    if pd.isna(r['Instrument Name']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Instrument Name'],
        'Quantity': r['Quantity'],
        'Market (Rs. cr)': r['Market Value'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── DSP ───────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'DSP_IRS_Output.xlsx', header=0)
df = df[~df['Fund Name'].astype(str).str.startswith('Red rows')]  # drop legend
for _, r in df.iterrows():
    if pd.isna(r['Name of Instrument']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Name of Instrument'],
        'Quantity': r['Quantity (Notional Rs.)'],
        'Market (Rs. cr)': r['Market Value (Rs. Cr)'],
    }])

# ── HDFC ──────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'HDFC_IRS_Output.xlsx', header=0)
for _, r in df.iterrows():
    scheme = r['Scheme Name (from file)']
    if pd.isna(scheme) or str(scheme).startswith('E.'):
        continue
    all_rows += make_rows([{
        'Fund Name': scheme,
        'Name of the Instrument': r['Underlying Security'],
        'Quantity': r['Quantity (Notional Value Rs. in lacs.)'],
        'Market (Rs. cr)': r['Market Value (Rs. in lacs.)'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── ICICI ─────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'ICICI_IRS_Output.xlsx', header=0)
for _, r in df.iterrows():
    inst = r['Company/Issuer/Instrument Name']
    if pd.isna(inst):
        continue
    if 'INTEREST RATE SWAPS' in str(inst).upper() and 'NOTIONAL' in str(inst).upper():
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': inst,
        'Quantity': r['Quantity'],
        'Market (Rs. cr)': r['Exposure/Market Value(Rs.Lakh)'],
        '% to Net\n Assets': r['% to Nav'],
        'YTM': r['Yield of the instrument'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── KOTAK ─────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'Kotak_IRS_Output.xlsx', header=1)
for _, r in df.iterrows():
    if pd.isna(r['Name of the Instrument']) or pd.isna(r['Market (Rs. cr)']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Name of the Instrument'],
        'Quantity': r['Notional (Rs.)'],
        'Market (Rs. cr)': r['Market (Rs. cr)'],
        'YTM': r['YTM'],
        'Maturity(days)': r['Residual Maturity (days)'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── NIPPON ────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'Nippon_IRS_Output.xlsx', header=0)
for _, r in df.iterrows():
    if pd.isna(r['Name of Instrument']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Name of Instrument'],
        'Quantity': r['Quantity (FV in Lacs)'],
        'Market (Rs. cr)': r['Market Value (Rs. Cr)'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── SBI ───────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'SBI_IRS_Output.xlsx', header=0)
for _, r in df.iterrows():
    if pd.isna(r['Instrument Name']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Instrument Name'],
        'Quantity': r['Quantity'],
        'Market (Rs. cr)': r['Market Value'],
        'Maturity Date (Rating)': r['Maturity Date'],
    }])

# ── ABSL ──────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'ABSL_IRS_Output.xlsx', header=1)
df = df[df['Fund Name'].notna() & ~df['Fund Name'].astype(str).str.startswith('▶')]
# Find maturity column flexibly — name varies across versions
absl_mat_col = next((c for c in df.columns if 'maturity date' in str(c).lower()), None)
for _, r in df.iterrows():
    if pd.isna(r['Underlying Security']):
        continue
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Underlying Security'],
        'Quantity': r['Notional (Rs. lakh)'],
        'Market (Rs. cr)': r['Market (Rs. cr)'],
        'YTM': r['YTM'],
        'Maturity(days)': r['Residual Maturity (days)'],
        'Maturity Date (Rating)': r[absl_mat_col] if absl_mat_col else None,
        'Swap Type': r['Swap Type'],
        'Long Position': r['Long Position'],
        'Short Positon': r['Short Position'],
    }])

# ── TATA ──────────────────────────────────────────────────────────────────────
df = pd.read_excel(INPUT_DIR + 'TATA_IRS_Output.xlsx', header=0)
df = df[df['Underlying'].notna()]
for _, r in df.iterrows():
    all_rows += make_rows([{
        'Fund Name': r['Fund Name'],
        'Name of the Instrument': r['Underlying'],
        'Quantity': r['Notional Value (in Lakhs)'],
        'Market (Rs. cr)': r['Market Value (Rs. Cr)'],
        'Maturity Date (Rating)': r['Maturity/Next Interest Fixing Date'],
    }])

# ── BUILD DATAFRAME ───────────────────────────────────────────────────────────
out = pd.DataFrame(all_rows, columns=OUT_COLS)
print(f"Total rows: {len(out)}")
print(out[['Fund Name', 'Market (Rs. cr)', 'Swap Type', 'Maturity Date']].head(20).to_string())

# ── WRITE EXCEL WITH FORMATTING ───────────────────────────────────────────────
out.to_excel(OUTPUT_PATH, index=False, sheet_name='Sheet1')

wb = load_workbook(OUTPUT_PATH)
ws = wb['Sheet1']

# Header style
hdr_fill = PatternFill('solid', start_color='1F4E79', end_color='1F4E79')
hdr_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for cell in ws[1]:
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = hdr_align
    cell.border = border

# Data rows alternating fill
fill_even = PatternFill('solid', start_color='EBF3FB', end_color='EBF3FB')
fill_odd = PatternFill('solid', start_color='FFFFFF', end_color='FFFFFF')
data_font = Font(name='Arial', size=9)
data_align_center = Alignment(horizontal='center', vertical='center')
data_align_left = Alignment(horizontal='left', vertical='center')

# Identify column indices for specific formatting
num_cols = {'Quantity', 'Market (Rs. cr)', '% to Net\n Assets', 'YTM',
            'Maturity(days)', 'Duration', 'sumproduct yield',
            'sumproducct Maturity', 'sumproduct Duration', 'Coupon'}
date_cols = {'Date', 'Maturity Date (Rating)', 'Maturity Date', 'Yr-Month'}

col_map = {cell.value: cell.column for cell in ws[1]}

for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
    fill = fill_even if row_idx % 2 == 0 else fill_odd
    for cell in row:
        col_name = ws.cell(row=1, column=cell.column).value
        cell.fill = fill
        cell.font = data_font
        cell.border = border
        if col_name in num_cols:
            cell.alignment = data_align_center
            if cell.value is not None and col_name in {'YTM', 'Coupon', '% to Net\n Assets'}:
                cell.number_format = '0.0000%' if col_name == 'YTM' else '0.00%'
            elif cell.value is not None and col_name == 'Market (Rs. cr)':
                cell.number_format = '#,##0.00'
        elif col_name in date_cols:
            cell.alignment = data_align_center
            if cell.value is not None:
                cell.number_format = 'DD-MMM-YYYY'
        else:
            cell.alignment = data_align_left

# Auto-fit column widths (capped)
for col in ws.columns:
    max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
    ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 10), 35)

ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

wb.save(OUTPUT_PATH)
print("Done. Saved to", OUTPUT_PATH)