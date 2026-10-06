import sys
import os
import argparse
from difflib import SequenceMatcher
from openpyxl import load_workbook, Workbook


def log(msg: str):
    print(f"[INFO] {msg}")


def similar(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a.strip().lower(), b.strip().lower()).ratio()


def is_blank(val) -> bool:
    return val is None or str(val).strip() == ""


def starts_with_IN(val) -> bool:
    return not is_blank(val) and str(val).strip().upper().startswith("IN")


def copy_cell_number_format(src, dest):
    dest.number_format = src.number_format


def copy_row_values_with_fund(src_sheet, dest_sheet, src_row, dest_row, fund_name: str):
    """
    Copy entire source row to output while inserting fund name between A and B.
    DO NOT move metadata placement elsewhere in the sheet.
    Output layout:
      A_out = A_src
      B_out = fund_name
      C_out = B_src
      D_out = C_src
      E_out = D_src
      ...
    """
    # A -> A
    srcA = src_sheet.cell(row=src_row, column=1)
    destA = dest_sheet.cell(row=dest_row, column=1)
    destA.value = srcA.value
    copy_cell_number_format(srcA, destA)

    # Inject fund name at B
    destB = dest_sheet.cell(row=dest_row, column=2)
    destB.value = fund_name
    destB.number_format = "General"

    # Copy source from B onward shifted by +1
    out_col = 3  # C
    for src_col in range(2, src_sheet.max_column + 1):
        src_cell = src_sheet.cell(row=src_row, column=src_col)
        dest_cell = dest_sheet.cell(row=dest_row, column=out_col)
        dest_cell.value = src_cell.value
        copy_cell_number_format(src_cell, dest_cell)
        out_col += 1


def row_has_grand_total(sheet, row_idx: int) -> bool:
    """Return True if any cell in this row contains 'Grand Total' (case/space-insensitive)."""
    for cell in sheet[row_idx]:
        if cell.value and "grand total" in str(cell.value).strip().lower():
            return True
    return False


def extract_fund_name(sheet) -> str:
    """Return first non-blank in row 1 (handles merged top-left)."""
    for cell in sheet[1]:
        if not is_blank(cell.value):
            name = str(cell.value).strip()
            log(f"Selected FUND NAME: '{name}'")
            return name
    log("WARNING: No non-blank fund name in row 1. Using empty string.")
    return ""


def main(input_file: str, output_file: str, fuzzy_threshold: float = 0.90):
    if not os.path.exists(input_file):
        print(f"[ERROR] Input file not found: {input_file}")
        sys.exit(1)

    log(f"Loading workbook: {input_file}")
    wb = load_workbook(input_file, data_only=False)
    out_wb = Workbook()
    out_sheet = out_wb.active
    out_sheet.title = "matched_rows"

    # These are checked in the BLANK-CHECK column (now column B)
    target_phrases = [
        "TREPS / Reverse Repo Investments",
        "Net Receivables/Payables "
    ]

    # ==== COLUMN MAP (UPDATED) ====
    ISIN_COL = 3          # Column C
    BLANK_CHECK_COL = 2   # Column B
    DATA_START_ROW = 5

    output_row = 1
    # Keep metadata placement exactly as before: based on (sheet.max_column - 1)
    max_copy_cols = 0

    log("Starting sheet processing...")

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        log(f"\nProcessing sheet: {sheet_name}")

        fund_name = extract_fund_name(sheet)

        # Preserve original metadata positioning rule
        copy_cols = sheet.max_column - 1  # (legacy: excluding Sr No)
        if copy_cols < 0:
            copy_cols = 0
        if copy_cols > max_copy_cols:
            max_copy_cols = copy_cols

        for row_idx in range(DATA_START_ROW, sheet.max_row + 1):

            if row_has_grand_total(sheet, row_idx):
                log(f"Grand Total encountered in '{sheet_name}' at row {row_idx}. Next sheet.")
                break

            isin_val = sheet.cell(row=row_idx, column=ISIN_COL).value

            # CASE 1: Valid ISIN in column C
            if not is_blank(isin_val) and starts_with_IN(isin_val):
                copy_row_values_with_fund(sheet, out_sheet, row_idx, output_row, fund_name)

                # Do NOT shift metadata (keep legacy placement)
                meta_start = max_copy_cols + 1
                out_sheet.cell(row=output_row, column=meta_start).value = sheet_name
                out_sheet.cell(row=output_row, column=meta_start + 1).value = fund_name
                out_sheet.cell(row=output_row, column=meta_start + 2).value = row_idx

                log(f"Copied row {row_idx} from '{sheet_name}' (valid ISIN: {isin_val})")
                output_row += 1
                continue

            # CASE 2: ISIN blank -> match phrases in column B
            if is_blank(isin_val):
                col_val = sheet.cell(row=row_idx, column=BLANK_CHECK_COL).value

                for phrase in target_phrases:
                    if similar(str(col_val), phrase) >= fuzzy_threshold:
                        copy_row_values_with_fund(sheet, out_sheet, row_idx, output_row, fund_name)

                        meta_start = max_copy_cols + 1  # unchanged
                        out_sheet.cell(row=output_row, column=meta_start).value = sheet_name
                        out_sheet.cell(row=output_row, column=meta_start + 1).value = fund_name
                        out_sheet.cell(row=output_row, column=meta_start + 2).value = row_idx

                        log(f"Copied row {row_idx} from '{sheet_name}' (blank ISIN matched '{phrase}')")
                        output_row += 1
                        break

    log(f"\nSaving output to: {output_file}")
    os.makedirs(os.path.dirname(output_file) or ".", exist_ok=True)
    out_wb.save(output_file)
    log(f"Extraction complete. Total rows copied: {output_row - 1}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="DSP extraction script (ISIN in C, checks in B) with fund-name insertion.")
    parser.add_argument("input_file", help="Path to the input Excel file")
    parser.add_argument("output_file", help="Path to the output Excel file")
    parser.add_argument("--fuzzy-threshold", type=float, default=0.90)
    args = parser.parse_args()
    main(args.input_file, args.output_file, args.fuzzy_threshold)
