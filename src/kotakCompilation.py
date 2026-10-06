import openpyxl
from openpyxl import Workbook
import re
import datetime

LOG_FILE = "bank8_log.txt"

TARGETS = {
    "isin": "isin",
    "triparty": "triparty repo",
    "net_assets": "net current assets/(liabilities)",
    "stop_key": "grand total"
}

def log(msg):
    t = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = msg.encode("ascii", "replace").decode("ascii")
    print(f"{t} {safe}")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{t} {safe}\n")

def clean(v):
    if v is None:
        return ""
    v = str(v).replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    v = re.sub(r"\s+", " ", v)
    return v.strip().lower()

# NEW helper
def is_triparty_or_ccil(text: str) -> bool:
    t = clean(text)
    return (
        "triparty repo" in t
        or "the clearing corporation of india limited" in t
    )

def append_with_header(ws, out_ws, r, header):
    row_data = []
    for c in range(1, 11):
        if c == 10:
            row_data.append(header)
        else:
            row_data.append(ws.cell(row=r, column=c).value)
    out_ws.append(row_data)

def process_sheet(ws, out_ws):

    header = ws["C1"].value
    log(f"Header extracted: {header}")

    max_row = ws.max_row
    r = 3

    while r <= max_row:

        colA = clean(ws[f"A{r}"].value)
        colC = clean(ws[f"C{r}"].value)
        colD = clean(ws[f"D{r}"].value)
        colE = clean(ws[f"E{r}"].value)

        if colE == TARGETS["stop_key"]:
            log(f"Row {r}: GRAND TOTAL found -> STOP")
            break

        if colD != "":
            if colD == TARGETS["isin"]:
                r += 1
                continue

            append_with_header(ws, out_ws, r, header)
            log(f"Row {r}: ISIN -> COPIED")
            r += 1
            continue

        # UPDATED RULE 2
        if colC != "" and is_triparty_or_ccil(colC):
            append_with_header(ws, out_ws, r, header)
            log(f"Row {r}: Triparty Repo / CCIL -> COPIED")
            r += 1
            continue

        if colC == "" and colA == TARGETS["net_assets"]:
            append_with_header(ws, out_ws, r, header)
            log(f"Row {r}: Net Current Assets/(Liabilities) -> COPIED")
            r += 1
            continue

        r += 1

def process_file(input_path, output_path):
    open(LOG_FILE, "w").close()
    log("=== BANK #8 SCRIPT STARTED ===")

    wb = openpyxl.load_workbook(input_path, data_only=True)
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    for sheet_name in wb.sheetnames:
        log(f"\n-- Processing Sheet: {sheet_name} --")
        ws = wb[sheet_name]
        process_sheet(ws, out_ws)

    out_wb.save(output_path)
    log(f"\nOutput saved: {output_path}")
    log("=== BANK #8 SCRIPT COMPLETE ===")

if __name__ == "__main__":
    process_file(
        input_path="Kotak.xlsx",
        output_path="./CompiledFiles/Compiled_Kotak.xlsx"
    )
