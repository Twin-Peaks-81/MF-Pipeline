import openpyxl
from openpyxl import Workbook
import os
import re
import datetime

LOG_FILE = "bank7_log.txt"

TARGETS = {
    "isin": "isin",
    # "treps" removed from exact match usage; we'll use contains check
    "treps_end": "sub total",
    "net_assets": "net receivables / (payables)",
    "stop_key": "grand total"
}

# =========================
# LOGGING
# =========================

def log(msg):
    t = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = msg.encode("ascii", "replace").decode("ascii")
    print(f"{t} {safe}")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{t} {safe}\n")

# =========================
# CLEAN FUNCTION
# =========================

def clean(v):
    if v is None:
        return ""
    v = str(v).lower()
    v = v.replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", v).strip()

# -------------------------
# TREPS detector (contains)
# -------------------------
def is_treps(text: str) -> bool:
    """Return True if cell text indicates the TREPS section header,
    ignoring leading '(a)/(b)/(c)' or other junk. Uses contains."""
    t = clean(text)
    return ("treps" in t) and ("reverse repo" in t)

# =========================
# PROCESS ONE FILE
# =========================

def process_single_file(filepath, out_ws):
    log(f"\n=== Processing file: {filepath} ===")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # Header from B1
    header = ws["B1"].value
    log(f"Header extracted: {header}")

    max_row = ws.max_row
    max_col = ws.max_column

    r = 9   # <<<<<< START ROW

    while r <= max_row:

        colB = clean(ws[f"B{r}"].value)
        colC = clean(ws[f"C{r}"].value)

        # Inject header into source col A (as per earlier spec)
        ws[f"A{r}"].value = header

        # STOP if GRAND TOTAL in column B
        if colB == TARGETS["stop_key"]:
            log(f"Row {r}: GRAND TOTAL → STOP FILE")
            break

        # ======================================
        # RULE 1: Column C HAS DATA (ISIN)
        # ======================================
        if colC != "":
            if colC == TARGETS["isin"]:
                log(f"Row {r}: C = ISIN header → SKIPPED")
                r += 1
                continue

            # Copy ISIN row
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, max_col + 1)
            ])
            log(f"Row {r}: ISIN row → COPIED")
            r += 1
            continue

        # ======================================
        # RULE 2: Column C blank → CHECK B
        # ======================================

        # Net Receivables / (Payables)
        if colB == TARGETS["net_assets"]:
            out_ws.append([
                ws.cell(row=r, column=c).value
                for c in range(1, max_col + 1)
            ])
            log(f"Row {r}: Net Receivables → COPIED")
            r += 1
            continue

        # TREPS BLOCK (contains match)
        if is_treps(colB):
            log(f"Row {r}: TREPS TRIGGER (contains) → BEGIN TREPS BLOCK")
            r += 1

            while r <= max_row:
                nextB = clean(ws[f"B{r}"].value)

                if nextB == TARGETS["treps_end"]:
                    log(f"Row {r}: Sub Total → END TREPS BLOCK")
                    break

                out_ws.append([
                    ws.cell(row=r, column=c).value
                    for c in range(1, max_col + 1)
                ])
                log(f"Row {r}: TREPS BLOCK → COPIED")
                r += 1

            continue

        # Anything else → SKIPPED
        log(f"Row {r}: No match → SKIPPED")
        r += 1


# =========================
# PROCESS FOLDER
# =========================

def process_folder(folder_path, output_file):
    open(LOG_FILE, "w").close()
    log("=== BANK #7 SCRIPT STARTED ===")

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    files_processed = 0

    for filename in os.listdir(folder_path):
        if filename.startswith("~$"):
            continue
        if filename.endswith(".xlsx"):
            process_single_file(os.path.join(folder_path, filename), out_ws)
            files_processed += 1

    out_wb.save(output_file)

    log("===== SCRIPT COMPLETED =====")
    log(f"TOTAL FILES PROCESSED: {files_processed}")
    log(f"Output saved: {output_file}")


# =========================
# RUN SCRIPT
# =========================

if __name__ == "__main__":
    process_folder(
        folder_path="Mirae",
        output_file="./CompiledFiles/Compiled_Mirae.xlsx"
    )
