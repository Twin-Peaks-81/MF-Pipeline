import openpyxl
from openpyxl import Workbook
from rapidfuzz import fuzz
import datetime
import re

TARGETS = {
    "grand_total": "grand total",
    "isin": "isin",
    "net_current_assets": "net current assets",
    "treps": "triparty repo/ reverse repo instrument",
    "cash_margin": "cash margin - ccil",
    "subtotal": "sub total"
}

STRICT_FUZZ = 99
SOFT_FUZZ = 95
LOG_FILE = "script_bank2_log.txt"


def log(msg):
    t = datetime.datetime.now().strftime("[%Y-%m-%d %H:%M:%S]")
    safe = msg.encode("ascii", "replace").decode("ascii")
    print(f"{t} {safe}")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"{t} {safe}\n")


def clean(v):
    if v is None:
        return ""
    v = str(v).lower()
    v = v.replace("\xa0", " ").replace("\n", " ").replace("\t", " ")
    v = re.sub(r"\s+", " ", v)
    return v.strip()


def exact_grand_total(v):
    return clean(v) == TARGETS["grand_total"]


def fuzzy_match(value, key):
    text = clean(value)
    target = TARGETS[key]
    threshold = SOFT_FUZZ if key in ("net_current_assets", "treps", "cash_margin") else STRICT_FUZZ
    return fuzz.partial_ratio(text, target) >= threshold


def process_file(input_path, output_path):
    open(LOG_FILE, "w").close()
    log("=== BANK #2 SCRIPT STARTED ===")

    wb = openpyxl.load_workbook(input_path, data_only=True)
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Compiled"

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        log(f"\n-- Processing sheet: {sheet_name} --")

        # Find GRAND TOTAL in column C
        grand_total_row = None
        for r in range(1, ws.max_row + 1):
            if exact_grand_total(ws[f"C{r}"].value):
                grand_total_row = r
                log(f"Found GRAND TOTAL at row {r}")
                break

        if not grand_total_row:
            log("Skipping sheet - no GRAND TOTAL")
            continue

        # Fill column A with B1
        header = ws["B1"].value
        for r in range(2, grand_total_row + 1):
            ws[f"A{r}"].value = header

        r = 2
        while r <= grand_total_row:
            colB = clean(ws[f"B{r}"].value)
            colC = clean(ws[f"C{r}"].value)

            # =============================
            # RULE 1: Column B has data
            # =============================
            if colB != "":

                # NEW RULE: If B starts with "p" → SKIP
                if colB.startswith("p"):
                    log(f"Row {r}: B starts with 'p' → SKIPPED")
                    r += 1
                    continue

                # If B contains ISIN → SKIP
                if fuzzy_match(colB, "isin"):
                    log(f"Row {r}: B has ISIN → SKIPPED")
                    r += 1
                    continue

                # Otherwise copy
                out_ws.append([ws.cell(row=r, column=c).value
                               for c in range(1, ws.max_column + 1)])
                log(f"Row {r}: B has data → COPIED")
                r += 1
                continue

            # =============================
            # RULE 2: B blank → check C
            # =============================

            if colC == "":
                log(f"Row {r}: B blank, C blank → SKIPPED")
                r += 1
                continue

            # C contains ISIN → skip
            if fuzzy_match(colC, "isin"):
                log(f"Row {r}: C contains ISIN → SKIPPED")
                r += 1
                continue

            # Cash Margin - CCIL
            if fuzzy_match(colC, "cash_margin"):
                out_ws.append([ws.cell(row=r, column=c).value
                               for c in range(1, ws.max_column + 1)])
                log(f"Row {r}: Cash Margin - CCIL → COPIED")
                r += 1
                continue

            # Net Current Assets
            if fuzzy_match(colC, "net_current_assets"):
                out_ws.append([ws.cell(row=r, column=c).value
                               for c in range(1, ws.max_column + 1)])
                log(f"Row {r}: Net Current Assets → COPIED")
                r += 1
                continue

            # TREPS block
            if fuzzy_match(colC, "treps"):
                log(f"Row {r}: TREPS trigger → BEGIN TREPS BLOCK")
                r += 1

                while r <= grand_total_row:
                    nextC = clean(ws[f"C{r}"].value)

                    if fuzzy_match(nextC, "subtotal"):
                        log(f"Row {r}: SUB TOTAL → END TREPS BLOCK")
                        break

                    out_ws.append([
                        ws.cell(row=r, column=c).value
                        for c in range(1, ws.max_column + 1)
                    ])
                    log(f"Row {r}: TREPS block row → COPIED")
                    r += 1

                continue

            # ANY OTHER DATA IN C → SKIP
            log(f"Row {r}: C has random data → SKIPPED")
            r += 1

    out_wb.save(output_path)
    log(f"\nSaved output to {output_path}")
    log("=== BANK #2 SCRIPT DONE ===")


if __name__ == "__main__":
    process_file(
        input_path="Nippon.xlsx",
        output_path="./CompiledFiles/Compiled_Nippon.xlsx"
    )
