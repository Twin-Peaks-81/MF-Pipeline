import re
import os
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def extract_maturity_date(text):
    # Match dates like 17-09-2026 or 17/09/2026
    match = re.search(r'(\d{2}[-/]\d{2}[-/]\d{4})', text)
    return match.group(1) if match else None

import os
SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
input_path  = os.path.join(SCRIPT_DIR, "SBI.xlsx")
output_path = os.path.join(SCRIPT_DIR, "SBI_IRS_Output.xlsx")

wb_in = load_workbook(input_path, data_only=True)

wb_out = Workbook()
ws_out = wb_out.active
ws_out.title = "IRS Data"

headers = ["Instrument Name", "Quantity", "Market Value", "Maturity Date", "Fund Name"]
ws_out.append(headers)

header_font = Font(bold=True, name='Arial')
header_fill = PatternFill("solid", start_color="D9E1F2")
for col in range(1, 6):
    cell = ws_out.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')

rows_added = 0

for sheet_name in wb_in.sheetnames:
    ws = wb_in[sheet_name]

    # Fund name from D3
    fund_name = ws.cell(row=3, column=4).value
    fund_name = str(fund_name).strip() if fund_name else sheet_name

    in_irs = False
    fund_rows = 0

    for row in ws.iter_rows():
        c_cell = row[2] if len(row) > 2 else None  # Column C
        c_val = str(c_cell.value).strip() if c_cell and c_cell.value else ''

        if c_val.upper() == 'INTEREST RATE SWAPS':
            in_irs = True
            continue

        if in_irs:
            if 'derivatives total' in c_val.lower() or 'derivative total' in c_val.lower():
                in_irs = False
                continue

            if not c_val:
                continue

            # Quantity from column H (index 7)
            h_cell = row[7] if len(row) > 7 else None
            try:
                quantity = float(h_cell.value)
            except (TypeError, ValueError):
                continue

            maturity = extract_maturity_date(c_val)
            market_value = round(quantity / 100, 4)

            ws_out.append([c_val, quantity, market_value, maturity, fund_name])
            rows_added += 1
            fund_rows += 1

    if fund_rows > 0:
        print(f"  [FOUND]  {fund_name} — {fund_rows} IRS row(s)")
    else:
        print(f"  [NONE]   {fund_name} — no IRS rows")

# Style data rows
for row in ws_out.iter_rows(min_row=2, max_row=ws_out.max_row):
    for cell in row:
        cell.font = Font(name='Arial')
        cell.alignment = Alignment(horizontal='left' if cell.column in (1, 5) else 'right')

ws_out.column_dimensions['A'].width = 60
ws_out.column_dimensions['B'].width = 18
ws_out.column_dimensions['C'].width = 18
ws_out.column_dimensions['D'].width = 15
ws_out.column_dimensions['E'].width = 50

wb_out.save(output_path)
print(f"\nDone. {rows_added} IRS rows written to {output_path}")
