import pandas as pd
import os
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Config ─────────────────────────────────────────────────────────────────────
SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE  = os.path.join(SCRIPT_DIR, "tata.xlsx")
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "TATA_IRS_Output.xlsx")

# ── Date input ──────────────────────────────────────────────────────────────────
date_input = "30-Apr-2026"
try:
    DATE_VAL = datetime.strptime(date_input, "%d-%b-%Y")
except ValueError:
    DATE_VAL = None

def header_matches(cell_text):
    """
    Match 'Hedging Positions through Interest Rate Swaps' in col B.
    If a valid date was entered, also check that the date appears in the string.
    Date matching is flexible (regex) so '15th April 2026', '15-Apr-2026', etc. all work.
    """
    text = str(cell_text).strip()
    if "interest rate swap" not in text.lower():
        return False
    if "hedging" not in text.lower():
        return False
    if DATE_VAL is None:
        return True
    # Build a loose pattern: day (with optional suffix), month name, year
    day   = str(DATE_VAL.day)
    month = DATE_VAL.strftime("%B")          # e.g. April
    mon3  = DATE_VAL.strftime("%b")          # e.g. Apr
    year  = str(DATE_VAL.year)
    pattern = rf"{day}(st|nd|rd|th)?\W*({month}|{mon3})\W*{year}"
    return bool(re.search(pattern, text, re.IGNORECASE))

# ── Output columns ──────────────────────────────────────────────────────────────
OUT_COLS = [
    "Fund Name",
    "Underlying",
    "Position",
    "Instrument Type",
    "Maturity/Next Interest Fixing Date",
    "Notional Value (in Lakhs)",
    "Market Value (Rs. Cr)",
]

# ── Styling helpers ──────────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")
ALT_FILL     = PatternFill("solid", fgColor="EBF3FB")
WHITE_FILL   = PatternFill("solid", fgColor="FFFFFF")
NOTE_FILL    = PatternFill("solid", fgColor="FFF2CC")

HEADER_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
SECTION_FONT = Font(name="Arial", bold=True, color="1F3864", size=9)
DATA_FONT    = Font(name="Arial", size=9)
NOTE_FONT    = Font(name="Arial", bold=True, size=10, color="7F6000")

thin    = Side(style="thin", color="BFBFBF")
BORDER  = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT    = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT   = Alignment(horizontal="right",  vertical="center")

def sc(cell, font=None, fill=None, align=None, border=None, nf=None):
    if font:   cell.font          = font
    if fill:   cell.fill          = fill
    if align:  cell.alignment     = align
    if border: cell.border        = border
    if nf:     cell.number_format = nf

# ── Extract IRS rows from one sheet ─────────────────────────────────────────────
def extract_irs(df):
    """
    Returns list of dicts with keys matching OUT_COLS (minus Fund Name),
    only for rows where Instrument Type == 'fixed' (case-insensitive).
    """
    header_row = None
    results    = []

    for i, row in df.iterrows():
        b_val = row.iloc[1] if len(row) > 1 else ""

        # Find the section header
        if header_row is None:
            if header_matches(b_val):
                header_row = i
            continue

        # We're inside the IRS section — next row is the column-header row, skip it
        if i == header_row + 1:
            continue

        # Stop at footnote / blank sentinel
        b_str = str(b_val).strip()
        if b_str.startswith("##") or b_str.startswith("^^"):
            break
        if b_str.upper() == "NIL":
            break
        all_blank = all(str(v) in ("nan", "", "None") for v in row.values)
        if all_blank:
            break

        # Column D (index 3) = Instrument Type
        inst_type = str(row.iloc[3]).strip() if len(row) > 3 else ""
        if inst_type.lower() != "fixed":
            continue

        underlying  = b_val if str(b_val) not in ("nan", "") else ""
        position    = str(row.iloc[2]).strip() if len(row) > 2 else ""
        mat_date    = row.iloc[4] if len(row) > 4 else ""
        notional    = row.iloc[5] if len(row) > 5 else ""

        # Market Value = Notional / 100
        try:
            market_val = float(notional) / 100
        except (TypeError, ValueError):
            market_val = ""

        results.append({
            "Underlying":                         underlying,
            "Position":                           position,
            "Instrument Type":                    inst_type,
            "Maturity/Next Interest Fixing Date": mat_date,
            "Notional Value (in Lakhs)":          notional,
            "Market Value (Rs. Cr)":              market_val,
        })

    return results

# ── Main loop ───────────────────────────────────────────────────────────────────
xl = pd.ExcelFile(INPUT_FILE)
sheets_to_process = [s for s in xl.sheet_names if s.lower() != "index"]

funds_with_irs    = []   # (sheet_name, [rows])
funds_without_irs = []

for sheet in sheets_to_process:
    df        = pd.read_excel(INPUT_FILE, sheet_name=sheet, header=None)
    # Read fund name from cell B1 (row 0, col 1); fall back to sheet name if blank
    b1        = df.iloc[0, 1] if df.shape[1] > 1 else ""
    fund_name = str(b1).strip() if str(b1) not in ("nan", "", "None") else sheet
    rows = extract_irs(df)
    if rows:
        funds_with_irs.append((fund_name, rows))
    else:
        funds_without_irs.append(fund_name)

# ── Build workbook ───────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "TATA IRS Fixed"

# Header row
ws.row_dimensions[1].height = 30
for col_idx, col_name in enumerate(OUT_COLS, start=1):
    c = ws.cell(row=1, column=col_idx, value=col_name)
    sc(c, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=BORDER)

current_row = 2

for fund_name, data_rows in funds_with_irs:

    # Section banner
    ws.row_dimensions[current_row].height = 18
    c = ws.cell(row=current_row, column=1, value=fund_name)
    sc(c, font=SECTION_FONT, fill=SECTION_FILL, align=LEFT, border=BORDER)
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row, end_column=len(OUT_COLS))
    current_row += 1

    for row_idx, rec in enumerate(data_rows):
        fill = WHITE_FILL if row_idx % 2 == 0 else ALT_FILL
        ws.row_dimensions[current_row].height = 15

        # Col 1: Fund Name
        c = ws.cell(row=current_row, column=1, value=fund_name)
        sc(c, font=DATA_FONT, fill=fill, align=LEFT, border=BORDER)

        # Col 2: Underlying
        c = ws.cell(row=current_row, column=2, value=rec["Underlying"])
        sc(c, font=DATA_FONT, fill=fill, align=LEFT, border=BORDER)

        # Col 3: Position
        c = ws.cell(row=current_row, column=3, value=rec["Position"])
        sc(c, font=DATA_FONT, fill=fill, align=CENTER, border=BORDER)

        # Col 4: Instrument Type
        c = ws.cell(row=current_row, column=4, value=rec["Instrument Type"])
        sc(c, font=DATA_FONT, fill=fill, align=CENTER, border=BORDER)

        # Col 5: Maturity Date
        mat = rec["Maturity/Next Interest Fixing Date"]
        c   = ws.cell(row=current_row, column=5, value=mat if str(mat) not in ("nan", "", "None") else None)
        sc(c, font=DATA_FONT, fill=fill, align=CENTER, border=BORDER, nf="DD-MMM-YYYY")

        # Col 6: Notional Value
        notional = rec["Notional Value (in Lakhs)"]
        c = ws.cell(row=current_row, column=6,
                    value=notional if str(notional) not in ("nan", "", "None") else None)
        sc(c, font=DATA_FONT, fill=fill, align=RIGHT, border=BORDER, nf="#,##0.00")

        # Col 7: Market Value
        mv = rec["Market Value (Rs. Cr)"]
        c  = ws.cell(row=current_row, column=7, value=mv if mv != "" else None)
        sc(c, font=DATA_FONT, fill=fill, align=RIGHT, border=BORDER, nf="#,##0.00")

        current_row += 1

    current_row += 1  # spacer

# ── No-IRS funds note ────────────────────────────────────────────────────────────
if funds_without_irs:
    current_row += 1
    c = ws.cell(row=current_row, column=1,
                value="SHEETS WITH NO FIXED-TYPE IRS DATA:")
    sc(c, font=NOTE_FONT, fill=NOTE_FILL, align=LEFT, border=BORDER)
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row, end_column=len(OUT_COLS))
    current_row += 1
    for fn in funds_without_irs:
        c = ws.cell(row=current_row, column=1, value=fn)
        sc(c, font=Font(name="Arial", size=9, italic=True, color="595959"),
           fill=PatternFill("solid", fgColor="FFFDF0"),
           align=LEFT, border=BORDER)
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=len(OUT_COLS))
        current_row += 1

# ── Column widths ─────────────────────────────────────────────────────────────────
col_widths = [20, 65, 12, 18, 28, 26, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
wb.save(OUTPUT_PATH)

print(f"\nDone. Output: {OUTPUT_PATH}")
print(f"Sheets with FIXED IRS rows : {len(funds_with_irs)}")
for fn, rows in funds_with_irs:
    print(f"  [{len(rows)} rows] {fn}")
print(f"Sheets with no FIXED IRS   : {len(funds_without_irs)}")
for fn in funds_without_irs:
    print(f"  - {fn}")